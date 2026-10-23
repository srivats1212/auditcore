"""Warning letters: the FDA Warning Letters listing, then the full letter
HTML for each row (PRD section 4, source 2).

The listing is paged from the page's datatables AJAX endpoint (see
``fda_tables``). Each row links to the letter page, whose URL ends in
``<slug>-<CMS number>-<MMDDYYYY>``; the CMS number is the letter's key.
"""

from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from urllib.parse import urljoin

from bs4 import BeautifulSoup
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import WarningLetter
from ..text.ocr_cleanup import clean_text
from .common import Fetcher, find_fei, get_or_create_facility, parse_date, pick
from .fda_tables import as_dicts, cell_href, cell_text, iter_pages
from .store import RawArchive, ingestion_run, record_error, store_document, tally

log = logging.getLogger(__name__)

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
WL_LIST_PATH = "/inspections-compliance-enforcement-and-criminal-investigations/compliance-actions-and-activities/warning-letters"
_CMS_FROM_URL = re.compile(r"-(\d{5,7})-\d{6,8}/?$")


@dataclass
class IndexRow:
    company_name: str | None
    letter_url: str | None
    cms_number: str | None
    issue_date: date | None
    posted_date: date | None
    issuing_office: str | None
    subject: str | None


# Fallback for slugs where the date is replaced by a site name:
# ".../kellin-chemicals-co-ltd-716228-zhangjiagang", ".../excelvision-726714-fareva".
_CMS_IN_SLUG = re.compile(r"-(\d{6})(?=-|/?$)")


def cms_from_url(url: str | None) -> str | None:
    if not url:
        return None
    path = url.split("?")[0]
    m = _CMS_FROM_URL.search(path) or _CMS_IN_SLUG.search(path.rstrip("/").rsplit("/", 1)[-1])
    return m.group(1) if m else None


def parse_index_xlsx(data: bytes, base_url: str = "https://www.fda.gov") -> list[IndexRow]:
    wb = load_workbook(io.BytesIO(data), read_only=False, data_only=True)
    ws = wb.active
    rows = ws.iter_rows()
    header_cells = next(rows)
    headers = [str(c.value).strip() if c.value is not None else "" for c in header_cells]
    out: list[IndexRow] = []
    for cells in rows:
        row = {h: c.value for h, c in zip(headers, cells)}
        if not any(v not in (None, "") for v in row.values()):
            continue
        url = pick(row, "Letter URL", "URL", "Link")
        if not url:
            # The export usually hyperlinks the company name cell.
            for h, c in zip(headers, cells):
                if c.hyperlink is not None and c.hyperlink.target:
                    url = c.hyperlink.target
                    break
        if url:
            url = urljoin(base_url, str(url).strip())
        cms = pick(row, "CMS Number", "MARCS-CMS Number", "CMS #", "CMS") or cms_from_url(url)
        out.append(
            IndexRow(
                company_name=pick(row, "Company Name", "Company", "Firm Name"),
                letter_url=url,
                cms_number=str(cms).strip() if cms else None,
                issue_date=parse_date(pick(row, "Letter Issue Date", "Issue Date", "Letter Date")),
                posted_date=parse_date(pick(row, "Posted Date", "Posted")),
                issuing_office=pick(row, "Issuing Office", "Office"),
                subject=pick(row, "Subject"),
            )
        )
    return out


def letter_text_from_html(html: str) -> str:
    soup = BeautifulSoup(html, "lxml")
    for tag in soup(["script", "style", "nav", "header", "footer", "aside", "form"]):
        tag.decompose()
    node = (
        soup.select_one("article")
        or soup.select_one("[role=main]")
        or soup.select_one("main")
        or soup.select_one("#main-content")
        or soup.body
        or soup
    )
    # Keep block structure so numbered violations start on their own line.
    for br in node.find_all("br"):
        br.replace_with("\n")
    blocks = [el.get_text(" ", strip=True) for el in node.find_all(["p", "li", "h1", "h2", "h3", "h4", "div"])
              if not el.find(["p", "li", "div"])]
    text = "\n\n".join(b for b in blocks if b)
    return text or node.get_text("\n", strip=True)


def _upsert_row(session: Session, row: IndexRow) -> WarningLetter:
    wl = session.scalars(select(WarningLetter).where(WarningLetter.cms_number == row.cms_number)).first()
    if wl is None:
        wl = WarningLetter(cms_number=row.cms_number)
        session.add(wl)
    wl.company_name = wl.company_name or row.company_name
    wl.issue_date = wl.issue_date or row.issue_date
    wl.posted_date = wl.posted_date or row.posted_date
    wl.issuing_office = wl.issuing_office or row.issuing_office
    wl.subject = wl.subject or row.subject
    wl.letter_url = wl.letter_url or row.letter_url
    session.flush()
    return wl


def _ingest_rows(session: Session, archive: RawArchive, fetcher: Fetcher | None, run, rows: list[IndexRow], *,
                 fetch_letters: bool, limit: int | None, subject: str | None, refetch: bool) -> dict:
    fetched = kept = 0
    for row in rows:
        if subject and subject.lower() not in (row.subject or "").lower():
            continue
        if not row.cms_number:
            record_error(run, row.company_name or "?", "no CMS number in index row")
            continue
        kept += 1
        wl = _upsert_row(session, row)
        if not (fetch_letters and fetcher and row.letter_url):
            continue
        if wl.document_id is not None and not refetch:
            continue  # incremental: already have it; --refetch re-checks for revised letters
        if limit is not None and fetched >= limit:
            continue
        try:
            resp = fetcher.get(row.letter_url)
            fetched += 1
            ingest_letter_html(session, archive, wl, resp.content, url=row.letter_url, run=run)
        except Exception as exc:  # one bad letter must not stop the run
            record_error(run, f"wl:{row.cms_number}", exc)
    return {"rows": len(rows), "matched": kept, "fetched": fetched, "errors": run.n_errors}


def ingest_index(session: Session, archive: RawArchive, fetcher: Fetcher | None, xlsx: bytes, *,
                 source_url: str | None = None, fetch_letters: bool = True, limit: int | None = None,
                 subject: str | None = None, refetch: bool = False) -> dict:
    """Offline path from an XLSX index that carries letter links. FDA's own
    export does not, so ``ingest_listing`` is the normal path."""
    with ingestion_run(session, "warning_letters") as run:
        idx = store_document(session, archive, source="wl_index", source_key="wl_index", data=xlsx,
                             media_type=XLSX, source_url=source_url, run=run)
        tally(run, idx.outcome)
        return _ingest_rows(session, archive, fetcher, run, parse_index_xlsx(xlsx), fetch_letters=fetch_letters,
                            limit=limit, subject=subject, refetch=refetch)


def rows_from_listing(cols: list[str], cells: list[list[str]], base_url: str) -> list[IndexRow]:
    out = []
    for row in as_dicts(cols, cells):
        url = cell_href(row.get("Company Name", ""), base_url)
        out.append(IndexRow(
            company_name=cell_text(row.get("Company Name", "")) or None,
            letter_url=url,
            cms_number=cms_from_url(url),
            issue_date=parse_date(cell_text(row.get("Letter Issue Date", ""))),
            posted_date=parse_date(cell_text(row.get("Posted Date", ""))),
            issuing_office=cell_text(row.get("Issuing Office", "")) or None,
            subject=cell_text(row.get("Subject", "")) or None,
        ))
    return out


def ingest_listing(session: Session, archive: RawArchive, fetcher: Fetcher, *, fetch_letters: bool = True,
                   limit: int | None = None, max_rows: int | None = None, subject: str | None = None,
                   refetch: bool = False) -> dict:
    """Page through the Warning Letters listing (newest first), archive each
    listing page, then fetch letters."""
    page_url = urljoin(fetcher.settings.fda_base_url, WL_LIST_PATH)
    totals = {"rows": 0, "matched": 0, "fetched": 0}
    with ingestion_run(session, "warning_letters") as run:
        rows: list[IndexRow] = []
        for start, (cols, raw, cells) in enumerate(iter_pages(fetcher, page_url, max_rows=max_rows)):
            page = store_document(session, archive, source="wl_index", source_key=f"wl_index:page:{start}",
                                  data=raw, media_type="application/json", source_url=page_url, run=run)
            tally(run, page.outcome)
            rows.extend(rows_from_listing(cols, cells, page_url))
        out = _ingest_rows(session, archive, fetcher, run, rows, fetch_letters=fetch_letters, limit=limit,
                           subject=subject, refetch=refetch)
        totals.update(out)
    return totals


_NOT_FOUND = re.compile(r"page you are looking for is not available|page not found", re.I)


class LetterNotAvailable(ValueError):
    """FDA serves its "page not available" page with HTTP 200 for withdrawn or moved letters."""


def is_not_found_page(html: bytes) -> bool:
    head = html[:200_000].decode("utf-8", errors="replace")
    title = re.search(r"<title>(.*?)</title>", head, re.I | re.S)
    return bool(_NOT_FOUND.search(head)) and not re.search(r"warning letter", title.group(1) if title else "", re.I)


def ingest_letter_html(session: Session, archive: RawArchive, wl: WarningLetter, html: bytes, *,
                       url: str | None, run=None) -> str:
    if is_not_found_page(html):
        raise LetterNotAvailable(f"{url or wl.cms_number}: FDA returned its not-available page")
    result = store_document(session, archive, source="warning_letter", source_key=f"wl:{wl.cms_number}",
                            data=html, media_type="text/html", source_url=url, run=run)
    if run is not None:
        tally(run, result.outcome)
    doc = result.document
    if result.outcome != "skipped" or doc.text_raw is None:
        doc.text_raw = letter_text_from_html(html.decode("utf-8", errors="replace"))
        doc.text_clean = clean_text(doc.text_raw)
    wl.document_id = doc.id if doc.is_current else wl.document_id
    fei = find_fei(doc.text_clean)
    if fei or wl.company_name:
        fac = get_or_create_facility(session, fei=fei, name=wl.company_name)
        wl.facility_id = fac.id if fac else wl.facility_id
    session.flush()
    return result.outcome


def ingest_letter_files(session: Session, archive: RawArchive, paths: list[Path]) -> dict:
    """Offline path: ingest saved letter HTML files named by CMS number
    (e.g. ``640123.html``) for letters already present in the index."""
    with ingestion_run(session, "warning_letters_files") as run:
        for path in paths:
            cms = path.stem
            wl = session.scalars(select(WarningLetter).where(WarningLetter.cms_number == cms)).first()
            if wl is None:
                wl = WarningLetter(cms_number=cms)
                session.add(wl)
                session.flush()
            try:
                ingest_letter_html(session, archive, wl, path.read_bytes(), url=wl.letter_url, run=run)
            except Exception as exc:
                record_error(run, str(path), exc)
        return {"files": len(paths), "errors": run.n_errors}
